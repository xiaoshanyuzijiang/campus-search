"""Document reading and original-download boundaries, with no external requests."""

import socket
import asyncio
import unittest
from io import BytesIO
from unittest.mock import AsyncMock, patch
from urllib.parse import parse_qs, urlsplit

import httpx
from docx import Document
from openpyxl import Workbook

from backend.tools.files import extract_document, fetch_public, make_download_url, response_filename, validate_public_url, verify_download_token
from backend.tools.page_reader import extract_page, read_notice_page
from backend.tools.research import ResearchTools
from backend.schemas import SchoolIdentity


class DocumentTests(unittest.TestCase):
    def test_docx_includes_table_and_paragraphs_in_order(self):
        doc = Document()
        doc.add_paragraph("奖学金办理要求")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "截止时间"
        table.cell(0, 1).text = "11 月 1 日"
        doc.add_paragraph("请提交成绩单")
        output = BytesIO()
        doc.save(output)
        text, status = extract_document(output.getvalue(), "申请表.docx")
        self.assertLess(text.index("办理要求"), text.index("截止时间"))
        self.assertLess(text.index("11 月 1 日"), text.index("成绩单"))
        self.assertEqual(status, "已读取")

    def test_xlsx_and_chinese_text_can_be_read(self):
        book = Workbook()
        book.active.append(["姓名", "所需材料"])
        book.active.append(["示例学生", "成绩单"])
        output = BytesIO()
        book.save(output)
        text, _ = extract_document(output.getvalue(), "名单.xlsx")
        self.assertIn("成绩单", text)
        text, _ = extract_document("请提交材料".encode("gb18030"), "说明.txt")
        self.assertIn("提交材料", text)

    def test_legacy_files_are_not_pretended_to_be_read(self):
        text, status = extract_document(b"legacy", "申请表.doc")
        self.assertEqual(text, "")
        self.assertIn("暂不支持", status)

    def test_pdf_text_is_read_and_image_only_pdf_is_reported(self):
        from backend.exports import result_pdf
        from tests.test_exports import sample_result
        from reportlab.pdfgen.canvas import Canvas
        text, status = extract_document(result_pdf(sample_result()), "通知.pdf")
        self.assertIn("奖学金", text)
        self.assertEqual(status, "已读取")
        stream = BytesIO()
        canvas = Canvas(stream)
        canvas.rect(10, 10, 200, 200)
        canvas.save()
        text, status = extract_document(stream.getvalue(), "扫描件.pdf")
        self.assertEqual(text, "")
        self.assertIn("扫描版", status)

    def test_page_discovers_official_relative_attachments_only(self):
        page = extract_page('<article><p>通知正文</p><a href="../../files/form.pdf">申请表</a><a href="https://evil.test/file.docx">伪造</a></article>'.encode(), "https://news.example.edu.cn/notice/1", "example.edu.cn")
        self.assertEqual(len(page["files"]), 1)
        self.assertEqual(page["files"][0]["url"], "https://news.example.edu.cn/files/form.pdf")
        self.assertIn("通知正文", page["content"])

    def test_signed_download_token_rejects_tampering_and_expiry(self):
        url = make_download_url("https://example.edu.cn/file.pdf", "申请表.pdf")
        token = parse_qs(urlsplit(url).query)["token"][0]
        self.assertEqual(verify_download_token(token)["filename"], "申请表.pdf")
        with self.assertRaises(ValueError):
            verify_download_token(token + "bad")
        with patch("backend.tools.files.time.time", return_value=10**12), self.assertRaises(ValueError):
            verify_download_token(token)

    def test_response_filename_preserves_chinese_download_name(self):
        headers = httpx.Headers({"content-disposition": "attachment; filename*=UTF-8''%E7%94%B3%E8%AF%B7%E8%A1%A8.docx"})
        self.assertEqual(response_filename(headers, "https://example.edu.cn/download.jsp"), "申请表.docx")


class NetworkBoundaryTests(unittest.IsolatedAsyncioTestCase):
    async def test_blocked_school_page_can_use_tavily_to_find_attachment(self):
        def handler(request):
            if request.method == "GET":
                return httpx.Response(403)
            return httpx.Response(200, json={"results": [{"url": "https://example.edu.cn/notice/1", "raw_content": "奖学金申请通知。 [申请表](../files/form.pdf)"}]})
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            with patch("backend.tools.files.validate_public_url", new_callable=AsyncMock):
                page = await read_notice_page(client, "test", "https://example.edu.cn/notice/1", "example.edu.cn")
        self.assertIn("奖学金", page["content"])
        self.assertEqual(page["files"][0]["url"], "https://example.edu.cn/files/form.pdf")

    async def test_search_budget_prevents_further_external_requests(self):
        requests = []
        def handler(request):
            requests.append(request)
            return httpx.Response(200, json={"results": []})
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            session = ResearchTools(client, "test", SchoolIdentity(name="示例大学"))
            search = session.as_tools()[0]
            for _ in range(20):
                result = await search.ainvoke({"query": "奖学金"})
        self.assertEqual(len(requests), 4)
        self.assertIn("error", result)

    async def test_private_dns_is_rejected(self):
        loop = asyncio.get_running_loop()
        with patch.object(loop, "getaddrinfo", new=AsyncMock(return_value=[(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", 80))])):
            with self.assertRaises(ValueError):
                await validate_public_url("http://school.example.edu.cn/file.pdf")

    async def test_redirect_is_checked_before_next_request(self):
        requests = []
        async def validate(url):
            if "127.0.0.1" in url:
                raise ValueError("private")
        def handler(request):
            requests.append(str(request.url))
            return httpx.Response(302, headers={"location": "http://127.0.0.1/private"})
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            with patch("backend.tools.files.validate_public_url", side_effect=validate), self.assertRaises(ValueError):
                await fetch_public(client, "https://example.edu.cn/file.pdf")
        self.assertEqual(len(requests), 1)

    async def test_unknown_source_number_does_not_make_network_request(self):
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: self.fail("Unexpected network access"))) as client:
            tools = ResearchTools(client, "test", SchoolIdentity(name="示例大学"))
            file_tool = next(tool for tool in tools.as_tools() if tool.name == "read_file")
            result = await file_tool.ainvoke({"source_number": 999})
            self.assertIn("error", result)


if __name__ == "__main__":
    unittest.main()

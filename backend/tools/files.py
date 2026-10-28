"""Bounded public HTTP reads, document text extraction, and signed downloads."""

import asyncio
import base64
import hashlib
import hmac
import ipaddress
import json
import os
import re
import secrets
import socket
import time
from io import BytesIO
from pathlib import Path
from urllib.parse import quote, unquote, urljoin, urlsplit
from zipfile import ZipFile

import httpx
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 48000
FILE_TYPES = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt", ".md", ".csv", ".zip", ".rar", ".7z"}
_SIGNING_KEY = secrets.token_bytes(32)


def _signing_key() -> bytes:
    return os.getenv("FILE_DOWNLOAD_SECRET", "").encode() or _SIGNING_KEY


def file_extension(url: str, title: str = "") -> str:
    for value in (unquote(urlsplit(url).path), unquote(urlsplit(url).query), title):
        match = re.search(r"\.(pdf|docx?|xlsx?|pptx?|txt|md|csv|zip|rar|7z)(?:$|[\s?&#])", value, re.I)
        if match:
            return "." + match.group(1).lower()
    return ""


def is_file_link(url: str, title: str = "") -> bool:
    return bool(file_extension(url, title) or re.search(r"(?:download\.(?:jsp|php|aspx?)|wbfileid=|/__local/)", url, re.I))


def safe_filename(name: str) -> str:
    return re.sub(r'[\x00-\x1f<>:"/\\|?*]', "_", name).strip(" .")[:180] or "attachment"


def response_filename(headers: httpx.Headers, url: str, title: str = "") -> str:
    disposition = headers.get("content-disposition", "")
    utf8 = re.search(r"filename\*=UTF-8''([^;]+)", disposition, re.I)
    plain = re.search(r'filename="?([^";]+)', disposition, re.I)
    name = unquote(utf8.group(1)) if utf8 else plain.group(1).strip() if plain else ""
    if not name:
        candidate = unquote(urlsplit(url).path.rsplit("/", 1)[-1])
        name = candidate if Path(candidate).suffix.lower() in FILE_TYPES else title or "attachment"
    ext = file_extension(url, name)
    if not ext:
        ext = {"application/pdf": ".pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx", "text/plain": ".txt"}.get(headers.get("content-type", "").split(";", 1)[0], "")
    if ext and not name.lower().endswith(ext):
        name += ext
    return safe_filename(name)


async def validate_public_url(url: str) -> None:
    try:
        parsed = urlsplit(url)
        host = (parsed.hostname or "").rstrip(".").lower()
        if parsed.scheme not in {"http", "https"} or not host or parsed.username or parsed.password:
            raise ValueError("文件地址不是有效的公开网页地址。")
        if host == "localhost" or host.endswith((".localhost", ".local")):
            raise ValueError("不允许读取本机或内网地址。")
        addresses = await asyncio.wait_for(asyncio.get_running_loop().getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM), 5)
        if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
            raise ValueError("不允许读取本机或内网地址。")
    except (OSError, asyncio.TimeoutError) as error:
        raise ValueError("无法解析文件所在网站。") from error


async def fetch_public(client: httpx.AsyncClient, url: str, *, limit: int = MAX_FILE_BYTES, referer: str = "") -> tuple[bytes, httpx.Headers, str]:
    headers = {"User-Agent": "Mozilla/5.0 (compatible; SchoolNotice/1.0)"}
    if referer:
        headers["Referer"] = referer
    for _ in range(5):
        await validate_public_url(url)
        async with client.stream("GET", url, headers=headers, follow_redirects=False, timeout=18) as response:
            if response.is_redirect:
                location = response.headers.get("location")
                if not location:
                    raise ValueError("网站返回了无效的跳转地址。")
                url = urljoin(url, location)
                continue
            response.raise_for_status()
            try:
                if int(response.headers.get("content-length", "0")) > limit:
                    raise ValueError("文件超过 12 MB 读取上限。")
            except ValueError as error:
                if "上限" in str(error):
                    raise
            data = bytearray()
            async for chunk in response.aiter_bytes():
                data.extend(chunk)
                if len(data) > limit:
                    raise ValueError("文件超过读取大小上限。")
            return bytes(data), response.headers, str(response.url)
    raise ValueError("网站跳转次数过多，无法读取。")


def _check_zip(data: bytes) -> None:
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 60 * 1024 * 1024:
            raise ValueError("文件解压后的内容过大，无法自动读取。")


def extract_document(data: bytes, filename: str) -> tuple[str, str]:
    ext = Path(filename).suffix.lower()
    truncated = False
    if ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            return "", "文件已加密，请下载后输入密码查看。"
        pages = []
        for page in reader.pages[:80]:
            pages.append(page.extract_text() or "")
            if sum(map(len, pages)) > MAX_TEXT_CHARS:
                break
        text = "\n\n".join(pages)
        truncated = len(reader.pages) > len(pages)
        if not text.strip():
            return "", "扫描版 PDF 没有可提取文字，请下载原文件查看。"
    elif ext == ".docx":
        _check_zip(data)
        document = Document(BytesIO(data))
        parts = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
            else:
                parts.append(block.text)
            if sum(map(len, parts)) > MAX_TEXT_CHARS:
                truncated = True
                break
        text = "\n".join(parts)
    elif ext == ".xlsx":
        _check_zip(data)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        parts = []
        try:
            for sheet in workbook.worksheets[:8]:
                parts.append("工作表：" + sheet.title)
                for row in sheet.iter_rows(max_row=500, max_col=40, values_only=True):
                    parts.append(" | ".join(str(value) if value is not None else "" for value in row))
                    if sum(map(len, parts)) > MAX_TEXT_CHARS:
                        break
                truncated = truncated or (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 40
                if sum(map(len, parts)) > MAX_TEXT_CHARS:
                    break
            truncated = truncated or len(workbook.worksheets) > 8
        finally:
            workbook.close()
        text = "\n".join(parts)
    elif ext in {".txt", ".md", ".csv"}:
        encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
        for encoding in encodings:
            try:
                text = data.decode(encoding)
                break
            except UnicodeError:
                continue
        else:
            return "", "无法识别文件文字编码，请下载原文件查看。"
    else:
        return "", "此格式暂不支持自动读取，可下载原文件查看。"
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
    truncated = truncated or len(text) > MAX_TEXT_CHARS
    return text[:MAX_TEXT_CHARS], "已读取（内容较长，仅整理部分内容）" if truncated else "已读取" if text else "未提取到文字，请下载原文件查看。"


def make_download_url(url: str, filename: str, source_url: str = "") -> str:
    payload = json.dumps({"url": url, "filename": safe_filename(filename), "referer": source_url, "expires": int(time.time()) + 21600}, ensure_ascii=False, separators=(",", ":")).encode()
    encoded = base64.urlsafe_b64encode(payload).rstrip(b"=")
    signature = hmac.new(_signing_key(), encoded, hashlib.sha256).hexdigest()
    return "/api/files/download?token=" + quote(encoded.decode() + "." + signature)


def verify_download_token(token: str) -> dict:
    try:
        if len(token) > 12000:
            raise ValueError
        encoded, signature = token.rsplit(".", 1)
        expected = hmac.new(_signing_key(), encoded.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            raise ValueError
        payload = json.loads(base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)))
        if payload["expires"] < time.time() or not all(isinstance(payload[key], str) for key in ("url", "filename", "referer")):
            raise ValueError
        return payload
    except (ValueError, KeyError, TypeError) as error:
        raise ValueError("下载链接已失效，请重新检索后下载。") from error
